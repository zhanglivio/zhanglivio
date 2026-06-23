#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把原始客户导入表清洗成规范 Excel。
   - 国家：原始列 → 税号前缀反推 → 人工修正(OVERRIDES)，并标注来源
   - 城市：大小写/空格规范化（首字母大写）
   - 大区：意大利按城市映射
   - 输出：exports/客户清单_clean.xlsx（客户清单 + 汇总 两个表）
"""
import openpyxl, re, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter, defaultdict
from xlsx_to_data import CITY_COORDS, IT_REGION, OVERRIDES, norm

SRC = "/root/.claude/uploads/2193f9bd-9279-5349-b56a-9dd264e1949d/a74d56a2-1782127989377193__.xlsx"

def titlecase(s):
    return re.sub(r"[A-Za-z]+", lambda m: m.group(0).capitalize(), str(s).strip()) if s else ""

def raw_country(r):
    c = r[18] or r[28]
    return (str(c).strip().upper(), "原始") if c else (None, None)

def vat_country(r):
    for i in (22, 20):
        v = r[i]
        if v:
            m = re.match(r"\s*([A-Za-z]{2})", str(v))
            if m:
                cc = m.group(1).upper()
                return "GR" if cc == "EL" else cc
    return None

def main():
    ws = openpyxl.load_workbook(SRC, read_only=True, data_only=True)["Sheet1"]
    rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if r and r[0] and str(r[0]).strip()]

    records = []
    for r in rows:
        name = str(r[0]).strip()
        co, src = raw_country(r)
        if not co:
            v = vat_country(r)
            if v: co, src = v, "税号反推"
        cy = (str(r[17]).strip() if r[17] else (str(r[27]).strip() if r[27] else None))
        ov = OVERRIDES.get(name)
        if ov:
            if ov.get("country"): co, src = ov["country"], "人工修正"
            if ov.get("city"): cy = ov["city"]
        city = titlecase(cy) if cy else ""
        region = ""
        lat = lng = ""
        if city:
            key = norm(city)
            if key in CITY_COORDS:
                lat, lng = CITY_COORDS[key]
            if co == "IT" and key in IT_REGION:
                region = IT_REGION[key]
        try: debt = round(float(r[5] or 0), 2)
        except: debt = 0.0
        try: open_debt = round(float(r[4] or 0), 2)
        except: open_debt = 0.0
        sales = (str(r[1]).strip() if r[1] else "")
        records.append({
            "抬头": name, "国家": co or "", "国家来源": src or "缺失",
            "大区": region, "城市": city, "纬度": lat, "经度": lng,
            "期初欠款": open_debt, "欠款总额": debt, "专属销售": sales,
            "上图": "是" if co else "否",
        })

    # 排序：有国家在前，按欠款降序
    records.sort(key=lambda x: (x["上图"] != "是", -x["欠款总额"]))

    wb = openpyxl.Workbook()
    ws1 = wb.active; ws1.title = "客户清单"
    headers = ["抬头","国家","国家来源","大区","城市","纬度","经度","期初欠款","欠款总额","专属销售","上图"]
    ws1.append(headers)

    head_fill = PatternFill("solid", fgColor="1A1A1A")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, len(headers)+1):
        cell = ws1.cell(1, c); cell.fill = head_fill; cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws1.freeze_panes = "A2"; ws1.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    debt_fill = PatternFill("solid", fgColor="F2F2F2")
    for rec in records:
        ws1.append([rec[h] for h in headers])
        row = ws1.max_row
        if rec["欠款总额"] and rec["欠款总额"] > 0:
            for c in (8, 9):
                ws1.cell(row, c).fill = debt_fill
                ws1.cell(row, c).font = Font(bold=True)
        if rec["上图"] == "否":
            ws1.cell(row, 2).font = Font(color="C00000")
    # 列宽
    widths = {"抬头":34,"国家":7,"国家来源":11,"大区":18,"城市":16,"纬度":9,"经度":9,"期初欠款":11,"欠款总额":12,"专属销售":12,"上图":6}
    for i, h in enumerate(headers, 1):
        ws1.column_dimensions[get_column_letter(i)].width = widths.get(h, 12)
    # 货币格式
    for row in range(2, ws1.max_row+1):
        for c in (8, 9):
            ws1.cell(row, c).number_format = '#,##0.00'

    # 汇总表
    ws2 = wb.create_sheet("汇总（按国家）")
    cc = Counter(); dd = defaultdict(float)
    for rec in records:
        k = rec["国家"] or "(缺失)"
        cc[k] += 1; dd[k] += rec["欠款总额"]
    ws2.append(["国家","客户数","欠款合计"])
    for c in range(1,4):
        cell = ws2.cell(1,c); cell.fill = head_fill; cell.font = head_font
        cell.alignment = Alignment(horizontal="center")
    for k, n in cc.most_common():
        ws2.append([k, n, round(dd[k],2)])
    ws2.column_dimensions["A"].width=12; ws2.column_dimensions["B"].width=10; ws2.column_dimensions["C"].width=14
    ws2.freeze_panes = "A2"
    for row in range(2, ws2.max_row+1):
        ws2.cell(row,3).number_format = '#,##0.00'
    # 合计行
    ws2.append(["合计", sum(cc.values()), round(sum(dd.values()),2)])
    for c in range(1,4): ws2.cell(ws2.max_row,c).font = Font(bold=True)
    ws2.cell(ws2.max_row,3).number_format = '#,##0.00'

    os.makedirs(os.path.join(os.path.dirname(__file__),"..","exports"), exist_ok=True)
    out = os.path.join(os.path.dirname(__file__),"..","exports","客户清单_clean.xlsx")
    wb.save(out)
    placed = sum(1 for r in records if r["上图"]=="是")
    print(f"已生成 {out}")
    print(f"总客户 {len(records)} | 有国家 {placed} | 缺国家 {len(records)-placed} | 总欠款 €{sum(r['欠款总额'] for r in records):,.2f}")

if __name__ == "__main__":
    main()
